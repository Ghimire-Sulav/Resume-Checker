from io import BytesIO

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from werkzeug.datastructures import FileStorage

import pytest

from services.resume_parser import ResumeParseError, parse_resume


def test_extract_txt():
    upload = FileStorage(stream=BytesIO(b"Python developer"), filename="resume.txt")
    assert parse_resume(upload) == "Python developer"


def test_extract_docx():
    document = Document()
    document.add_paragraph("Education: Example University")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python"
    output = BytesIO()
    document.save(output)
    output.seek(0)
    text = parse_resume(FileStorage(stream=output, filename="resume.docx"))
    assert "Example University" in text
    assert "Python" in text


def test_extract_pdf():
    writer = PdfWriter()
    page = writer.add_blank_page(width=100, height=100)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({
            NameObject("/F1"): writer._add_object(font),
        }),
    })
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 10 50 Td (PDF resume text) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    output = BytesIO()
    writer.write(output)
    output.seek(0)
    assert "PDF resume text" in parse_resume(FileStorage(stream=output, filename="resume.pdf"))


def test_reject_unsupported_file():
    upload = FileStorage(stream=BytesIO(b"anything"), filename="resume.exe")
    with pytest.raises(ResumeParseError, match="PDF, DOCX, or TXT"):
        parse_resume(upload)
