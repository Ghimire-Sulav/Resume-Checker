import re
import unicodedata
from collections import defaultdict
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from werkzeug.datastructures import FileStorage


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_RESUME_CHARS = 30_000

SECTION_ALIASES = {
    "education": (
        "education", "academic background", "academic history", "academics",
        "educational background", "qualifications",
    ),
    "experience": (
        "experience", "work experience", "professional experience",
        "employment history", "employment", "internship experience",
        "internship", "internships", "work history", "professional history",
    ),
    "projects": ("projects", "personal projects", "academic projects", "project experience"),
    "technical_skills": (
        "technical skills", "technical skill", "skills", "core technical skills",
        "technical skills and tools", "technical skills & tools",
        "skills and technologies", "tools and technologies", "technologies",
        "technology stack", "technical expertise", "technical tools",
    ),
    "soft_skills": ("soft skills", "interpersonal skills", "core competencies"),
    "certifications": (
        "certifications", "certification", "certificates", "licenses",
        "licenses and certifications",
    ),
    "career_objective": ("career objective", "objective", "career profile"),
    "summary": ("professional summary", "profile summary", "summary", "profile"),
    "achievements": ("achievements", "awards", "honors", "accomplishments"),
    "leadership": ("leadership", "volunteer experience", "extracurricular activities"),
    "publications": ("publications", "research", "languages", "references"),
}
_ALIAS_LOOKUP = {
    re.sub(r"\s+", " ", alias.casefold()): section
    for section, aliases in SECTION_ALIASES.items()
    for alias in aliases
}
_ALIASES_BY_LENGTH = sorted(_ALIAS_LOOKUP, key=len, reverse=True)


class ResumeParseError(ValueError):
    """Raised when an uploaded resume cannot be safely extracted."""


def parse_resume(upload: FileStorage) -> str:
    filename = upload.filename or ""
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ResumeParseError("Upload a PDF, DOCX, or TXT file.")

    try:
        data = upload.read()
        if not data:
            raise ResumeParseError("The uploaded file is empty.")
        if suffix == ".txt":
            text = data.decode("utf-8-sig")
        elif suffix == ".pdf":
            reader = PdfReader(BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise ResumeParseError("Password-protected PDFs are not supported.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            document = Document(BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                parts.extend(
                    " | ".join(cell.text for cell in row.cells)
                    for row in table.rows
                )
            text = "\n".join(parts)
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError(
            "We couldn't read that file. Check that it is a valid, uncorrupted document."
        ) from exc

    text = text.replace("\x00", "").strip()
    if not text:
        raise ResumeParseError("No readable text was found in the resume.")
    if len(text) > MAX_RESUME_CHARS:
        raise ResumeParseError(
            f"Resume text is too long (maximum {MAX_RESUME_CHARS:,} characters)."
        )
    return text


def extract_resume_sections(text: str) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = defaultdict(list)
    current_section: str | None = None

    for raw_line in text.splitlines():
        line = unicodedata.normalize("NFKC", raw_line).strip()
        if not line:
            continue
        line = re.sub(r"^\s*[#>*•●▪◦‣\-–—]+\s*", "", line)
        line = re.sub(r"^\s*\d{1,2}\s*[.)]\s*", "", line)
        line = line.strip().strip("[]")
        if not line:
            continue

        heading = _recognize_heading(line)
        if heading:
            current_section, inline_content = heading
            if inline_content:
                sections[current_section].append(_clean_content(inline_content))
            continue

        if current_section:
            content = _clean_content(line)
            if content:
                sections[current_section].append(content)

    return {name: values[:40] for name, values in sections.items()}


def _recognize_heading(line: str) -> tuple[str, str] | None:
    normalized = re.sub(r"\s+", " ", line.casefold()).strip()
    normalized = normalized.strip(" :|—–-")
    if normalized in _ALIAS_LOOKUP:
        return _ALIAS_LOOKUP[normalized], ""

    for alias in _ALIASES_BY_LENGTH:
        if not normalized.startswith(alias):
            continue
        remainder = line[len(alias):].strip()
        if not remainder:
            return _ALIAS_LOOKUP[alias], ""
        if remainder[0] in ":|—–-":
            inline = remainder.lstrip(" :|—–-").strip()
            return _ALIAS_LOOKUP[alias], inline
        if remainder.casefold().startswith((" & ", " and ", " / ")):
            tail = re.sub(r"^(?:&|and|/)\s*", "", remainder, flags=re.IGNORECASE)
            other = _recognize_heading(tail)
            if other and not other[1]:
                return _ALIAS_LOOKUP[alias], ""
    return None


def _clean_content(line: str) -> str:
    line = re.sub(r"^\s*(?:[-*•●▪◦‣]|\d+[.)])\s*", "", line).strip()
    return line.strip("|").strip()
